"""
Research helper — dari casebook → rencana riset + keyword + query siap-klik + workbook CASE_RESEARCH.xlsx.

    python tools/research_helper.py --case ../case                         # pakai case_config.json
    python tools/research_helper.py --case ../case --casebook ../case/casebook.pdf
    python tools/research_helper.py --make-template tools/CASE_RESEARCH_TEMPLATE.xlsx

Output di folder case:
    RESEARCH_PLAN.html / .md   ← checklist ekstraksi casebook, keyword, query (Google Scholar, Crossref, Semantic Scholar, Garuda,
                                 OJK/AAJI/AAUI/BPS, Google Books) — disesuaikan jenis asuransi, perusahaan & driver hasil fast run
    CASE_RESEARCH.xlsx         ← workbook untuk hasil temuan; sheet case_facts sudah diisi kandidat kalimat berangka dari casebook
                                 (kolom use = N; ubah ke Y untuk dipakai). Notebook membaca file ini otomatis.
Jalankan ulang setelah fast run → query driver-spesifik ikut dibuat (membaca outputs/churn_insurance/tables/driver_evidence_matrix.csv).
"""
import argparse
import collections
import datetime as dt
import html
import json
import os
import re
import sys
import urllib.parse

for _st in (sys.stdout, sys.stderr):
    try:
        _st.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

SECTIONS = ["executive_summary", "introduction", "data", "methodology", "findings_drivers", "findings_timing", "findings_segments",
            "findings_model", "findings_explain", "business", "causal", "personas", "recommendations", "discussion", "limitations", "conclusion"]
THEMES = ["lifecycle", "payment", "price", "service", "affordability", "relationship", "channel", "engagement", "demographic",
          "geography", "general", "methods"]

STOP = set("""a an the and or of to in on for with by from at as is are was were be been this that these those it its their our we you they
he she his her them which who whom whose what when where why how not no yes into over under than then so such can could may might will would
shall should do does did done have has had having also more most less least very per each other any all some both either neither about across
dan atau yang di ke dari untuk dengan pada adalah ini itu akan oleh dalam sebagai juga tidak bahwa serta karena maka agar bagi para setiap
lebih kurang telah sudah belum harus dapat bisa tersebut antara hingga sampai jika bila namun tetapi saat ketika secara kami kita mereka
data case team teams participant participants peserta tim soal file files column columns kolom csv train test customer customers row rows
submission sample target id meaning""".split())

THEME_KW = {
    "lifecycle": r"tenure|new_customer|orig|start|since|lama|masa|age_of_policy|policy_age|duration|vintage",
    "payment": r"payment|bayar|debit|renew|late|overdue|arrear|frequency|grace|tunggak",
    "price": r"premium_change|increase|price|kenaikan|rate_change|discount|premium_to_income",
    "service": r"claim|klaim|complain|keluhan|settle|reject|service|ticket|nps|satisf",
    "affordability": r"income|salary|gaji|pendapatan|afford|credit|market_value|wealth",
    "relationship": r"product|policies|bundle|rider|household|n_policies",
    "channel": r"channel|agent|agen|broker|bank|online|source",
    "engagement": r"engag|digital|app|login|visit|usage|active",
    "demographic": r"\bage|usia|umur|gender|sex|marital|children|dependent|education|occupation",
    "geography": r"region|city|county|province|provinsi|state|kota|wilayah|lat|long|zip|postal",
}

THEME_QUERIES = {
    "lifecycle": ["early lapse insurance contract age", "\"first year\" lapse life insurance persistency", "policy duration churn survival analysis insurance",
                  "lapse polis tahun pertama asuransi"],
    "payment": ["payment method lapse life insurance", "automatic payment customer retention insurance", "premium payment frequency lapse determinants",
                "forgetting to pay lapse insurance", "metode pembayaran premi lapse"],
    "price": ["price elasticity insurance renewal retention", "premium increase churn motor insurance", "renewal pricing fairness price walking insurance",
              "kenaikan premi perpanjangan polis"],
    "service": ["claims satisfaction customer retention insurance", "complaint handling customer churn financial services",
                "claim settlement time policyholder satisfaction", "kepuasan klaim retensi nasabah asuransi"],
    "affordability": ["emergency fund hypothesis lapse", "income shock life insurance lapse", "premium affordability policy lapse household"],
    "relationship": ["cross-buying customer retention insurance", "multi-product household defection insurance", "bundling switching costs insurance"],
    "channel": ["distribution channel persistency insurance agent", "bancassurance persistency lapse", "agent turnover policy lapse"],
    "engagement": ["digital engagement customer churn insurance", "mobile app usage retention financial services"],
    "demographic": ["policyholder age gender lapse determinants", "life stage insurance lapse behavior"],
    "geography": ["regional differences insurance lapse", "spatial variation insurance churn"],
    "general": ["customer churn insurance machine learning", "lapse risk life insurance determinants review", "customer retention insurance strategy"],
    "methods": ["uplift modeling customer retention insurance", "expected maximum profit customer churn", "survival analysis customer churn insurance",
                "doubly robust estimation marketing intervention", "SHAP churn prediction interpretability", "explainable boosting machine insurance"],
}


def urls(q, kind="scholar"):
    e = urllib.parse.quote_plus(q)
    return {
        "scholar": f"https://scholar.google.com/scholar?q={e}&as_ylo=2012",
        "crossref": f"https://search.crossref.org/search/works?q={e}&from_ui=yes",
        "semantic": f"https://www.semanticscholar.org/search?q={e}",
        "garuda": f"https://garuda.kemdikbud.go.id/documents?search={e}",
        "google": f"https://www.google.com/search?q={e}",
        "books": f"https://www.google.com/search?tbm=bks&q={e}",
    }[kind]


def read_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        try:
            import fitz
            with fitz.open(path) as d:
                return "\n".join(p.get_text() for p in d)
        except ImportError:
            print("PyMuPDF tidak ada → python -m pip install PyMuPDF")
            return ""
    if ext == ".docx":
        try:
            from docx import Document
            return "\n".join(p.text for p in Document(path).paragraphs)
        except ImportError:
            return ""
    return open(path, encoding="utf-8", errors="replace").read()


def keywords(text, k=25):
    words = [w for w in re.findall(r"[A-Za-zÀ-ÿ][A-Za-zÀ-ÿ\-]{2,}", text.lower()) if w not in STOP]
    uni = collections.Counter(words)
    bi = collections.Counter(" ".join(p) for p in zip(words, words[1:]) if p[0] != p[1])
    tri = collections.Counter(" ".join(p) for p in zip(words, words[1:], words[2:]))
    cands = [(t, c * 3) for t, c in tri.items() if c >= 2] + [(t, c * 2) for t, c in bi.items() if c >= 2] + [(t, c) for t, c in uni.items() if c >= 3]
    out, seen = [], set()
    for t, _ in sorted(cands, key=lambda x: -x[1]):
        if any(t in s or s in t for s in seen):
            continue
        seen.add(t)
        out.append(t)
        if len(out) >= k:
            break
    return out


def fact_sentences(text, k=40):
    """Kalimat yang memuat angka/persen/uang/tahun → kandidat fakta case."""
    sents = re.split(r"(?<=[.!?])\s+|\n{2,}", re.sub(r"[ \t]+", " ", text))
    out = []
    for s in sents:
        s = s.strip().replace("\n", " ")
        if re.search(r"(\.csv|`|\|)", s) or s.startswith("-"):
            continue                                   # baris kamus data / daftar file, bukan fakta bisnis
        if 30 <= len(s) <= 400 and re.search(r"(\d+[.,]?\d*\s?%|rp\.?\s?\d|idr|usd|\$\s?\d|\b(19|20)\d{2}\b|\d+[.,]\d+|\b\d{2,}\b)", s, re.I):
            out.append(s)
    return out[:k]


def capitalized_terms(text, k=15):
    terms = re.findall(r"\b(?:PT|CV|Tbk)\.?\s+[A-Z][A-Za-z]+(?:\s+[A-Z][A-Za-z]+){0,4}|\b[A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,3}\b", text)
    c = collections.Counter(t.strip() for t in terms if len(t) > 4)
    return [t for t, n in c.most_common(k) if n >= 2]


def read_cfg(case):
    p = os.path.join(case, "case_config.json")
    if not os.path.exists(p):
        return {}
    txt = "\n".join(l for l in open(p, encoding="utf-8").read().splitlines() if not l.strip().startswith("//"))
    return json.loads(txt) if txt.strip() else {}


def drivers_from_outputs(case, cfg):
    out = cfg.get("OUTPUT_DIR", "./outputs/churn_insurance")
    out = out if os.path.isabs(out) else os.path.join(case, out)
    p = os.path.join(out, "tables", "driver_evidence_matrix.csv")
    if not os.path.exists(p):
        return []
    import pandas as pd
    ev = pd.read_csv(p)
    return ev["feature"].astype(str).head(10).tolist()


def theme_of(feature):
    f = feature.lower()
    for th, pat in THEME_KW.items():
        if re.search(pat, f):
            return th
    return "general"


def make_template(path, case_facts=None, labels=None):
    """Workbook CASE_RESEARCH.xlsx dengan dropdown section/tema."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    rows = [
        ["CASE RESEARCH WORKBOOK — diisi tim, dibaca otomatis oleh notebook (Section 24.3)"],
        [""],
        ["Sheet", "Isi", "Masuk ke laporan"],
        ["case_facts", "Fakta dari casebook (angka, tujuan manajemen, produk). Kandidat sudah diisi otomatis dari casebook; ubah use = Y untuk dipakai.",
         "Disisipkan verbatim di section yang dipilih (default: introduction)"],
        ["industry_facts", "Fakta pasar/regulasi dari OJK, AAJI, AAUI, BPS, berita resmi. Kalimat EN + sitasi (Penulis/Judul, Tahun) + APA lengkap.",
         "Introduction (atau section lain) + daftar pustaka otomatis"],
        ["literature", "Temuan jurnal/buku: 1 kalimat EN yang memuat sitasi + APA lengkap + tema (lifecycle/payment/price/...).",
         "Discussion: dipilih otomatis bila temanya muncul di driver hasil data (tema 'general' selalu boleh) + daftar pustaka"],
        ["custom_paragraphs", "Paragraf tulisan tim untuk section tertentu (start/end).", "Disisipkan persis di posisi itu"],
        ["feature_labels", "Nama kolom → label bisnis (mis. 'curr_ann_amt' → 'annual premium').", "Semua teks, tabel & figure laporan"],
        [""],
        ["Section yang valid:", ", ".join(SECTIONS)],
        ["Tema literatur:", ", ".join(THEMES)],
        ["Aturan:", "Tulis dalam bahasa artikel (EN). Setiap fakta eksternal WAJIB punya sitasi + APA; hanya pakai sumber yang sudah kamu buka & cek."],
        ["Format sitasi:", "Dalam teks: (Eling & Kochanski, 2013) / Hein et al. (2020) / (\"Judul berita singkat,\" 2025). APA: lihat RESEARCH_KIT.md §5."],
    ]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=13)
    for c in ws[3]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 95
    ws.column_dimensions["C"].width = 70
    specs = {
        "case_facts": (["use", "section", "statement_EN", "source_in_casebook", "note"], [16, 22, 100, 24, 30]),
        "industry_facts": (["use", "section", "statement_EN", "in_text_citation", "apa_reference", "url", "accessed", "verified"], [8, 18, 90, 30, 90, 40, 12, 10]),
        "literature": (["use", "theme", "finding_EN", "in_text_citation", "apa_reference", "doi_or_url", "verified"], [8, 16, 90, 30, 90, 40, 10]),
        "custom_paragraphs": (["use", "section", "position", "text"], [8, 22, 10, 140]),
        "feature_labels": (["column", "label"], [40, 60]),
    }
    hdr_fill = PatternFill("solid", fgColor="D9E2F3")
    for name, (cols, widths) in specs.items():
        s = wb.create_sheet(name)
        s.append(cols)
        for j, (c, w) in enumerate(zip(cols, widths)):
            cell = s.cell(row=1, column=j + 1)
            cell.font = Font(bold=True)
            cell.fill = hdr_fill
            s.column_dimensions[cell.column_letter].width = w
        s.freeze_panes = "A2"
        if "use" in cols:
            dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
            s.add_data_validation(dv)
            dv.add(f"A2:A500")
        if "section" in cols:
            dv = DataValidation(type="list", formula1='"' + ",".join(SECTIONS) + '"', allow_blank=True)
            s.add_data_validation(dv)
            col = chr(ord("A") + cols.index("section"))
            dv.add(f"{col}2:{col}500")
        if "theme" in cols:
            dv = DataValidation(type="list", formula1='"' + ",".join(THEMES) + '"', allow_blank=True)
            s.add_data_validation(dv)
            col = chr(ord("A") + cols.index("theme"))
            dv.add(f"{col}2:{col}500")
        if "position" in cols:
            dv = DataValidation(type="list", formula1='"start,end"', allow_blank=True)
            s.add_data_validation(dv)
            dv.add("C2:C500")
        for row in s.iter_rows(min_row=1, max_row=1):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    ex = wb["industry_facts"]
    ex.append(["N", "introduction", "Indonesia's insurance penetration was 2.84% at the end of 2024, which OJK described as largely stagnant over two decades "
               "(\"Penetrasi asuransi syariah,\" 2025).", "(\"Penetrasi asuransi syariah,\" 2025)",
               "Penetrasi asuransi syariah di Indonesia masih rendah, ini penyebabnya. (2025, March 18). Kontan.co.id. "
               "https://keuangan.kontan.co.id/news/penetrasi-asuransi-syariah-di-indonesia-masih-rendah-ini-penyebabnya",
               "https://keuangan.kontan.co.id/news/penetrasi-asuransi-syariah-di-indonesia-masih-rendah-ini-penyebabnya", "2026-10", "Y"])
    lit = wb["literature"]
    lit.append(["N", "payment", "Forgetting to pay explained 37.8% of recent life-insurance lapses (Gottlieb & Smetters, 2021).", "(Gottlieb & Smetters, 2021)",
                "Gottlieb, D., & Smetters, K. (2021). Lapse-based insurance. American Economic Review, 111(8), 2377–2416. https://doi.org/10.1257/aer.20160868",
                "https://doi.org/10.1257/aer.20160868", "Y"])
    cp = wb["custom_paragraphs"]
    cp.append(["N", "recommendations", "end", "Contoh: Given the company's plan to launch its mobile app in Q1 (casebook p. 4), the first three actions should be "
               "delivered through in-app nudges before agent calls."])
    cf = wb["case_facts"]
    for sfx in (case_facts or []):
        cf.append(["N", "introduction", sfx, "", "auto-extracted — cek & ubah use=Y kalau dipakai"])
    fl = wb["feature_labels"]
    for a, b in (labels or []):
        fl.append([a, b])
    wb.save(path)
    return path


def build_plan(case, cfg, text, drivers):
    company = cfg.get("COMPANY_NAME") or "the company"
    line = cfg.get("INSURANCE_LINE", "auto")
    market = cfg.get("MARKET", "Indonesia")
    year = dt.date.today().year
    kws = keywords(text) if text else []
    terms = capitalized_terms(text) if text else []
    facts = fact_sentences(text) if text else []
    line_terms = {"life": ["life insurance", "asuransi jiwa", "lapse", "surrender", "unit link"],
                  "general": ["general insurance", "asuransi umum", "non-renewal", "motor insurance", "asuransi kendaraan"],
                  "health": ["health insurance", "asuransi kesehatan", "non-renewal"], "multi": ["insurance", "asuransi", "lapse", "non-renewal"]}
    lt = line_terms.get(line, ["insurance", "asuransi", "churn", "lapse"])
    Q = collections.OrderedDict()
    if str(market).lower().startswith("indo"):
        Q["A. Konteks industri Indonesia (regulator & asosiasi)"] = [
            (f"statistik perasuransian OJK {year - 1}", "google", "site:ojk.go.id"),
            (f"OJK penetrasi densitas asuransi {year - 1}", "google", ""),
            (f"AAJI kinerja asuransi jiwa {year - 1} klaim surrender", "google", ""),
            (f"AAUI kinerja asuransi umum {year - 1} premi", "google", ""),
            (f"SNLIK {year - 1} indeks literasi inklusi perasuransian", "google", ""),
            ("Program Penjaminan Polis LPS", "google", ""),
            (f"BPS statistik asuransi {year - 1}", "google", "site:bps.go.id"),
            ("POJK perlindungan konsumen dan masyarakat sektor jasa keuangan", "google", "site:ojk.go.id"),
        ]
    else:
        Q["A. Industry context"] = [(f"{lt[0]} lapse rate statistics {year - 1}", "google", ""), (f"{lt[0]} customer retention study {year - 1}", "google", ""),
                                     ("J.D. Power insurance shopping study", "google", "")]
    Q["B. Perusahaan / case"] = [(f"\"{company}\" annual report {year - 1}", "google", ""), (f"\"{company}\" laporan tahunan", "google", ""),
                                 (f"\"{company}\" persistency OR lapse OR retensi", "google", "")] + [(f"\"{t}\" {lt[1]}", "google", "") for t in terms[:4]]
    th_seen = []
    for f in drivers:
        th = theme_of(f)
        if th not in th_seen:
            th_seen.append(th)
    if not th_seen:
        th_seen = ["lifecycle", "payment", "price", "service"]
    Q["C. Literatur per driver (dari hasil fast run)" if drivers else "C. Literatur per driver (umum — jalankan ulang setelah fast run)"] = [
        (q, "scholar", th) for th in th_seen for q in THEME_QUERIES.get(th, [])[:3]]
    Q["D. Literatur Indonesia (jurnal nasional)"] = [(q, "garuda", "") for q in ["lapse polis asuransi", "churn nasabah asuransi", "retensi nasabah asuransi",
                                                                                "persistensi polis asuransi jiwa"]]
    Q["E. Metode (untuk menjawab juri & memperkuat Appendix B)"] = [(q, "scholar", "methods") for q in THEME_QUERIES["methods"]]
    Q["F. Keyword dari casebook (otomatis)"] = [(f"{k} {lt[0]}", "scholar", "") for k in kws[:10]]
    return Q, kws, terms, facts, th_seen


def render(case, cfg, Q, kws, terms, facts, themes, drivers):
    team = cfg.get("TEAM_NAME", "Team")
    md = [f"# RESEARCH PLAN — {team}", "", f"Dibuat {dt.datetime.now():%Y-%m-%d %H:%M}. Ikuti urutan: **1 → 2 → 3 → 4**. Target waktu total: 60–90 menit (bisa paralel dengan full run).", "",
          "## 1. Ekstraksi casebook (15 menit) → `case_config.json` + sheet `case_facts`", "",
          "| Cari di casebook | Isi ke |", "|---|---|",
          "| Nama perusahaan, jenis produk/lini, pasar | `COMPANY_NAME`, `INSURANCE_LINE`, `MARKET` |",
          "| Masalah & tujuan manajemen (2–4 kalimat) | `CASE_CONTEXT`, `MANAGEMENT_OBJECTIVE` |",
          "| Definisi churn/lapse/non-renewal, periode data | `CHURN_DEFINITION`, `DATA_PERIOD` |",
          "| Pertanyaan/tugas yang diminta | `CASE_QUESTIONS` (satu per baris) |",
          "| Intervensi yang mungkin (diskon, agen, app, auto-debit) | `AVAILABLE_INTERVENTIONS` |",
          "| Angka bisnis (margin, biaya kampanye, success rate, budget) | parameter bisnis + `BUSINESS_PARAMS_SOURCE = case` |",
          "| Kamus data (arti kolom) | sheet `feature_labels` |",
          "| Angka/fakta penting lain (pertumbuhan, pangsa pasar, rencana) | sheet `case_facts` (use = Y) |", ""]
    if terms or kws:
        md += ["**Istilah kunci terdeteksi di casebook:** " + ", ".join(terms[:12] + kws[:12]), ""]
    if facts:
        md += [f"**{len(facts)} kalimat berangka diambil otomatis** ke sheet `case_facts` (use = N). Contoh:", ""] + [f"- {f}" for f in facts[:8]] + [""]
    md += ["## 2. Cari referensi (30–45 menit) — klik query, buka sumber, catat ke `CASE_RESEARCH.xlsx`", "",
           "Prioritas sumber: **casebook > regulator (OJK, BI, LPS, BPS) > asosiasi (AAJI, AAUI) > jurnal peer-review (DOI) > buku > laporan praktisi (Bain, McKinsey, Swiss Re, RGA) > berita yang mengutip rilis resmi.**", ""]
    for sec, qs in Q.items():
        md += [f"### {sec}", "", "| Query | Scholar | Crossref | Semantic | Garuda | Google | Books |", "|---|---|---|---|---|---|---|"]
        for q, kind, tag in qs:
            qq = (q + " " + tag) if tag.startswith("site:") else q
            md.append(f"| {qq}{' — *' + tag + '*' if tag and not tag.startswith('site:') else ''} | [▶]({urls(qq, 'scholar')}) | [▶]({urls(qq, 'crossref')}) | "
                      f"[▶]({urls(qq, 'semantic')}) | [▶]({urls(qq, 'garuda')}) | [▶]({urls(qq, 'google')}) | [▶]({urls(qq, 'books')}) |")
        md.append("")
    md += ["## 3. Verifikasi (wajib, 2 menit per sumber)", "",
           "- [ ] Sumber benar-benar dibuka (bukan hanya cuplikan Google). Angka dicek di halaman aslinya.",
           "- [ ] Jurnal: DOI dibuka di https://doi.org/ → judul, penulis, tahun, volume, halaman cocok. Hindari jurnal predator (cek SINTA/Scopus).",
           "- [ ] Berita: pastikan mengutip rilis resmi (OJK/AAJI/AAUI) + tanggal; tulis 'menurut OJK' bila angka dari regulator.",
           "- [ ] Tulis 1 kalimat EN yang SUDAH berisi sitasi, mis. *Hein et al. (2020) found that payment frequency …*",
           "- [ ] Isi APA 7 lengkap (lihat contoh di `RESEARCH_KIT.md` §5). Kolom `verified` = Y.", "",
           "## 4. Masukkan ke laporan (otomatis)", "",
           "1. Simpan `CASE_RESEARCH.xlsx` di folder case (sudah dibuat).",
           "2. Jalankan ulang notebook **cell 24.2c–24.5 saja** (atau RUN_FULL). Fakta & literatur `use = Y` otomatis masuk ke section yang dipilih + daftar pustaka.",
           "3. Cek `REPORT_TODO.md` → bagian 'Research yang dipakai' menampilkan apa saja yang masuk.", "",
           f"Tema driver terdeteksi: **{', '.join(themes)}**" + (f" (driver: {', '.join(drivers[:6])})" if drivers else ""), ""]
    md_txt = "\n".join(md)
    with open(os.path.join(case, "RESEARCH_PLAN.md"), "w", encoding="utf-8") as f:
        f.write(md_txt)
    body = []
    for line in md:
        line_h = html.escape(line)
        line_h = re.sub(r"\[▶\]\((.+?)\)", lambda m: f'<a href="{html.unescape(m.group(1))}" target="_blank" rel="noopener">▶</a>', line_h)
        line_h = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", line_h)
        body.append(line_h)
    out, in_table = [], False
    for line in body:
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(set(c) <= set("-") for c in cells):
                continue
            tag = "th" if not in_table else "td"
            if not in_table:
                out.append("<table>")
                in_table = True
            out.append("<tr>" + "".join(f"<{tag}>{c}</{tag}>" for c in cells) + "</tr>")
            continue
        if in_table:
            out.append("</table>")
            in_table = False
        if line.startswith("### "):
            out.append(f"<h3>{line[4:]}</h3>")
        elif line.startswith("## "):
            out.append(f"<h2>{line[3:]}</h2>")
        elif line.startswith("# "):
            out.append(f"<h1>{line[2:]}</h1>")
        elif line.startswith("- "):
            out.append(f"<li>{line[2:]}</li>")
        elif line:
            out.append(f"<p>{line}</p>")
    if in_table:
        out.append("</table>")
    page = ("<!doctype html><html lang='id'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
            "<title>Research Plan</title><style>body{font:15px/1.5 system-ui,Segoe UI,sans-serif;max-width:1100px;margin:0 auto;padding:16px;"
            "background:#fcfcfb;color:#111}table{border-collapse:collapse;width:100%;font-size:13px;margin:6px 0 16px}td,th{border-bottom:1px solid #ddd;"
            "padding:5px 8px;text-align:left}th{background:#eef3fb}a{color:#2a78d6;text-decoration:none;font-weight:600}h2{margin-top:28px}"
            "@media(prefers-color-scheme:dark){body{background:#1a1a19;color:#eee}th{background:#24324a}td,th{border-color:#333}}</style></head><body>"
            + "\n".join(out) + "</body></html>")
    with open(os.path.join(case, "RESEARCH_PLAN.html"), "w", encoding="utf-8") as f:
        f.write(page)


def run(case, casebook=None, quiet=False):
    case = os.path.abspath(case)
    cfg = read_cfg(case)
    text = ""
    cands = [casebook] if casebook else [os.path.join(case, n) for n in os.listdir(case)
                                          if re.search(r"(casebook|case_book|soal|brief|case)\.(pdf|docx|md|txt)$", n, re.I)]
    for c in cands:
        if c and os.path.exists(c):
            text += "\n" + read_text(c)
            if not quiet:
                print(f"Casebook dibaca: {c} ({len(text):,} karakter)")
    text += "\n" + str(cfg.get("CASE_CONTEXT", "")) + "\n" + "\n".join(cfg.get("CASE_QUESTIONS", []) or [])
    drivers = drivers_from_outputs(case, cfg)
    Q, kws, terms, facts, themes = build_plan(case, cfg, text, drivers)
    render(case, cfg, Q, kws, terms, facts, themes, drivers)
    xl = os.path.join(case, "CASE_RESEARCH.xlsx")
    if not os.path.exists(xl):
        make_template(xl, case_facts=facts)
        if not quiet:
            print(f"Workbook dibuat: {xl} ({len(facts)} kandidat fakta dari casebook)")
    elif facts:
        try:
            from openpyxl import load_workbook
            wb = load_workbook(xl)
            ws = wb["case_facts"]
            have = {str(r[2].value) for r in ws.iter_rows(min_row=2) if r[2].value}
            added = 0
            for f_ in facts:
                if f_ not in have:
                    ws.append(["N", "introduction", f_, "", "auto-extracted"])
                    added += 1
            wb.save(xl)
            if added and not quiet:
                print(f"{added} kandidat fakta baru ditambahkan ke {xl}")
        except PermissionError:
            print("⚠️ CASE_RESEARCH.xlsx sedang dibuka di Excel → tutup dulu lalu jalankan ulang untuk menambah kandidat fakta.")
    if not quiet:
        print(f"✅ {os.path.join(case, 'RESEARCH_PLAN.html')} (buka di browser) & RESEARCH_PLAN.md")
    return os.path.join(case, "RESEARCH_PLAN.html")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--case", default=".")
    ap.add_argument("--casebook", default=None, help="PDF/DOCX/TXT/MD soal (default: cari file bernama casebook/soal/brief/case di folder case)")
    ap.add_argument("--make-template", default=None, help="hanya buat workbook kosong di path ini")
    a = ap.parse_args()
    if a.make_template:
        print("✅", make_template(a.make_template))
        return
    run(a.case, a.casebook)


if __name__ == "__main__":
    main()
